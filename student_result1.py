from openpyxl import Workbook, load_workbook
import os

FILE_NAME = "student_results.xlsx"


# -------------------------------------------------
# Create Excel file if it does not already exist
# -------------------------------------------------
def create_excel_file():

    if not os.path.exists(FILE_NAME):

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Student Results"

        headings = [
            "Roll No",
            "Name",
            "Class",
            "Subject 1",
            "Subject 2",
            "Subject 3",
            "Subject 4",
            "Subject 5",
            "Total",
            "Percentage",
            "Grade",
            "Status"
        ]

        sheet.append(headings)

        workbook.save(FILE_NAME)


# -------------------------------------------------
# Calculate total, percentage, grade and status
# -------------------------------------------------
def calculate_result(marks):

    total = sum(marks)

    percentage = total / 5

    # Check whether student passed every subject
    if all(mark >= 40 for mark in marks):
        status = "PASS"
    else:
        status = "FAIL"

    # Grade calculation
    if status == "FAIL":
        grade = "F"

    elif percentage >= 90:
        grade = "A+"

    elif percentage >= 80:
        grade = "A"

    elif percentage >= 70:
        grade = "B"

    elif percentage >= 60:
        grade = "C"

    elif percentage >= 50:
        grade = "D"

    else:
        grade = "E"

    return total, percentage, grade, status


# -------------------------------------------------
# Check if Roll Number already exists
# -------------------------------------------------
def roll_number_exists(roll_no):

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == str(roll_no):
            return True

    return False


# -------------------------------------------------
# Add Student Result
# -------------------------------------------------
def add_student():

    print("\n--------------------------------")
    print("       ADD STUDENT RESULT")
    print("--------------------------------")

    roll_no = input("Enter Student Roll No.: ")

    # Duplicate roll number check
    if roll_number_exists(roll_no):
        print("\nStudent with this Roll No. already exists.")
        return

    name = input("Enter Student Name: ")
    student_class = input("Enter Course/Class: ")

    marks = []

    print("\nEnter marks of 5 subjects:")

    for i in range(1, 6):

        while True:

            try:
                mark = float(input(f"Enter Subject {i} Marks: "))

                if 0 <= mark <= 100:
                    marks.append(mark)
                    break

                else:
                    print("Marks must be between 0 and 100.")

            except ValueError:
                print("Please enter valid numeric marks.")

    # Calculate result
    total, percentage, grade, status = calculate_result(marks)

    # Open Excel file
    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    # Add record
    sheet.append([
        roll_no,
        name,
        student_class,
        marks[0],
        marks[1],
        marks[2],
        marks[3],
        marks[4],
        total,
        round(percentage, 2),
        grade,
        status
    ])

    workbook.save(FILE_NAME)

    print("\n--------------------------------")
    print("Student Result Added Successfully")
    print("--------------------------------")

    print(f"Total       : {total}")
    print(f"Percentage  : {percentage:.2f}%")
    print(f"Grade       : {grade}")
    print(f"Status      : {status}")


# -------------------------------------------------
# Get Student Result
# -------------------------------------------------
def get_result():

    print("\n--------------------------------")
    print("        GET STUDENT RESULT")
    print("--------------------------------")

    roll_no = input("Enter Student Roll No.: ")

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    student_found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if str(row[0]) == str(roll_no):

            student_found = True

            print("\n--------------------------------")
            print("          Student Result")
            print("--------------------------------")

            print(f"Roll No     : {row[0]}")
            print(f"Name        : {row[1]}")
            print(f"Class       : {row[2]}")
            print(f"Total       : {row[8]}")
            print(f"Percentage  : {row[9]:.2f}%")
            print(f"Grade       : {row[10]}")
            print(f"Status      : {row[11]}")

            print("--------------------------------")

            break

    if not student_found:
        print("\nStudent record not found.")


# -------------------------------------------------
# Show All Student Data
# -------------------------------------------------
def show_all_data():

    workbook = load_workbook(FILE_NAME)
    sheet = workbook.active

    print("\n")
    print("=" * 85)
    print("                         ALL STUDENT RESULTS")
    print("=" * 85)

    print(
        f"{'Roll No':<10}"
        f"{'Name':<20}"
        f"{'Class':<12}"
        f"{'Total':<10}"
        f"{'Percentage':<15}"
        f"{'Grade':<10}"
        f"{'Status':<10}"
    )

    print("-" * 85)

    record_found = False

    for row in sheet.iter_rows(min_row=2, values_only=True):

        record_found = True

        percentage = f"{row[9]:.2f}%"

        print(
            f"{str(row[0]):<10}"
            f"{str(row[1]):<20}"
            f"{str(row[2]):<12}"
            f"{str(row[8]):<10}"
            f"{percentage:<15}"
            f"{str(row[10]):<10}"
            f"{str(row[11]):<10}"
        )

    if not record_found:
        print("No student records available.")

    print("=" * 85)


# -------------------------------------------------
# Main Menu
# -------------------------------------------------
def menu():

    create_excel_file()

    while True:

        print("\n========================================")
        print("      STUDENT RESULT MANAGEMENT")
        print("========================================")

        print("1. Add Student Result")
        print("2. Get Student Result")
        print("3. Show All Student Data")
        print("4. Exit")

        print("========================================")

        choice = input("Enter your choice: ")

        if choice == "1":
            add_student()

        elif choice == "2":
            get_result()

        elif choice == "3":
            show_all_data()

        elif choice == "4":
            print("\nThank you for using Student Result Management System.")
            print("Program Closed.")
            break

        else:
            print("\nInvalid choice. Please enter 1, 2, 3 or 4.")


# -------------------------------------------------
# Start Program
# -------------------------------------------------
menu()